#!/usr/bin/env python3
"""KOS voce 24 — confronto tra l'esportazione di Google Business Profile (file per gruppo) e i dati Yext delle schede
(strutture-yext-*.json). SOLA LETTURA. Abbinamento per sito web del profilo = uri della scheda.

Uso: python3 scripts/yext/confronto-gbp.py <cartella-gbp> <strutture-yext.json> <uscita.xlsx>
"""
import sys, re, json, glob, os, unicodedata
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

GIORNI = ['Mo', 'Tu', 'We', 'Th', 'Fr', 'Sa', 'Su']
COL_G = {'Mo': 'Orario del lunedì', 'Tu': 'Orario del martedì', 'We': 'Orario del mercoledì', 'Th': 'Orario del giovedì',
         'Fr': 'Orario del venerdì', 'Sa': 'Orario del sabato', 'Su': 'Orario della domenica'}

def norm(s):
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().lower()
    s = re.sub(r'\b(v|v\.le|vle)\b\.?', 'via', s)
    s = re.sub(r'\bp\.?za\b', 'piazza', s); s = re.sub(r'\bc\.?so\b', 'corso', s)
    return re.sub(r'[^a-z0-9]', '', s)

def tel(s):
    d = re.sub(r'\D', '', str(s or ''))
    return d[2:] if d.startswith('39') and d[2:3] in ('0', '3') and len(d) >= 10 else d

def orari_yext(s):
    out = {g: '' for g in GIORNI}
    for parte in re.split(r'[\n;]+', s or ''):
        m = re.match(r'\s*([A-Z][a-z])(?:-([A-Z][a-z]))?\s+(.+)', parte)
        if not m: continue
        a, b = GIORNI.index(m[1]), GIORNI.index(m[2] or m[1])
        for i in range(a, b + 1): out[GIORNI[i]] = m[3].replace(' ', '').replace('23:59', '24:00')
    return out

def orari_gbp(d):
    return {g: str(d.get(COL_G[g]) or '').replace(' ', '').replace('x', '') for g in GIORNI}

def fmt(o):
    return ' · '.join(f'{g} {o[g] or "chiuso"}' for g in GIORNI)

def main(cartella, yext_json, uscita):
    yext = {r['uri'].strip('/'): r['tv'] for r in json.load(open(yext_json))['risorse']}
    righe = []
    for f in sorted(glob.glob(os.path.join(cartella, '*.xlsx'))):
        gruppo = os.path.basename(f).split('-')[0]
        ws = openpyxl.load_workbook(f, read_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True)); h = rows[0]
        for r in rows[1:]:
            d = dict(zip(h, r)); d['_gruppo'] = gruppo
            righe.append(d)
    out, visti = [], set()
    for d in righe:
        sito = str(d.get('Sito web') or '')
        m = re.search(r'kosgroup\.com/(it/[^?#]+?)/?$', sito)
        uri = m[1] if m else ''
        y = yext.get(uri)
        diff = []
        if y:
            visti.add(uri)
            ind_g = ' '.join(str(d.get(c) or '') for c in ['Riga indirizzo 1', 'Riga indirizzo 2'])
            if norm(ind_g) != norm(y.get('str.via')): diff.append(('Indirizzo', ind_g.strip(), y.get('str.via')))
            if str(d.get('Codice postale') or '') != (y.get('str.cap') or ''): diff.append(('CAP', d.get('Codice postale'), y.get('str.cap')))
            if norm(d.get('Località')) != norm(y.get('str.citta')): diff.append(('Località', d.get('Località'), y.get('str.citta')))
            if tel(d.get('Telefono principale')) != tel(y.get('str.telefono')): diff.append(('Telefono', d.get('Telefono principale'), y.get('str.telefono')))
            og, oy = orari_gbp(d), orari_yext(y.get('str.orari'))
            if og != oy: diff.append(('Orari', fmt(og), fmt(oy)))
            if norm(d.get("Nome dell'attività commerciale")) != norm(y.get('str.nome')): diff.append(('Nome', d.get("Nome dell'attività commerciale"), y.get('str.nome')))
        if d['_gruppo'].startswith('Anni'): esito = 'Anni Azzurri (fuori voce 24)'
        elif not uri: esito = 'Sito non kosgroup.com/it'
        elif not y: esito = 'Scheda non in Yext/voce 24'
        elif diff: esito = 'Differenze'
        else: esito = 'Coincide'
        base = [d['_gruppo'], d.get('Stato'), d.get('Codice negozio'), d.get("Nome dell'attività commerciale"), sito,
                d.get('Categoria principale'), d.get('Categorie aggiuntive'), esito]
        if diff:
            for c, g, yv in diff: out.append(base + [c, g, yv])
        else: out.append(base + ['', '', ''])
    for uri in sorted(set(yext) - visti):
        out.append(['', '', '', yext[uri].get('str.nome'), 'https://kosgroup.com/' + uri, '', '', 'Scheda senza profilo GBP abbinato', '', '', ''])
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Confronto'
    ws.append(['Gruppo GBP', 'Stato', 'Codice negozio', 'Nome GBP', 'Sito web GBP', 'Categoria principale', 'Categorie aggiuntive',
               'Esito', 'Campo', 'Valore GBP', 'Valore Yext'])
    for r in out: ws.append([str(x) if x is not None else '' for x in r])
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E79')
    for col, w in zip('ABCDEFGHIJK', [16, 12, 16, 40, 50, 28, 30, 30, 11, 60, 60]): ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    wb.save(uscita)
    from collections import Counter
    prof = {}
    for r in out: prof[(r[3], r[4])] = r[7]
    print(Counter(prof.values()))
    print(Counter(r[8] for r in out if r[8]))

if __name__ == '__main__':
    main(*sys.argv[1:4])
