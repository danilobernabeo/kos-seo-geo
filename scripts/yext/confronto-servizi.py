#!/usr/bin/env python3
"""
Voce 24 — confronto delle associazioni struttura ↔ servizi e cure tra il sito kosgroup.com e i filtri «Servizi e Cure»
di Yext, per la validazione di KOS.

Uso:
    python3 scripts/yext/confronto-servizi.py <mappatura-servizi.json> <export-yext.xlsx> [--uscita kosgroup/dati/voce-24]

<mappatura-servizi.json> = output di scripts/checks/kos-mappatura-servizi.php (produzione). Si confrontano tre fonti:
    A  scheda → servizio   TV map* della scheda struttura (mapAree*, mapInterventi*)
    B  servizio → scheda   TV mapServiziStrutture della pagina del servizio
    C  Yext                colonna «Servizi e Cure» dell'entità abbinata alla scheda (dall'URL del sito web)
Solo le pagine della sezione «Servizi e cure» (convenzioni e tecnologie restano fuori). Le voci Yext si abbinano alle
pagine per nome, con gli alias di ALIAS per le voci scritte diversamente.
Esito: confronto-servizi-<data>.xlsx con il foglio «Confronto» (una riga per struttura e servizio, colonne Esito e Note
per KOS), «Riepilogo» (conteggi con formule) e «Legenda».
"""
import argparse
import collections
import json
import os
import re
import unicodedata
import warnings

ALIAS = {   # voce Yext → titolo della pagina del sito
    'MOC': 'Mineralometria Ossea Computerizzata - MOC',
    'TAC': 'Tomografia Assiale Computerizzata - TAC',
    'Maxillofacciale': 'Chirurgia maxillofacciale',
    'Urologia-Andrologia': 'Urologia e Andrologia',
    'Endrocrinologia': 'Endocrinologia',
    'Gastroentereologia': 'Gastroenterologia',
}
SEZIONI = {'ambulatori-e-medicina-specialistica': 'Ambulatori e medicina specialistica', 'psichiatria': 'Psichiatria',
           'riabilitazione': 'Riabilitazione'}


def norma(s):
    s = unicodedata.normalize('NFKD', s or '').encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()


def main():
    a = argparse.ArgumentParser()
    a.add_argument('mappatura')
    a.add_argument('export')
    a.add_argument('--uscita', default='kosgroup/dati/voce-24')
    o = a.parse_args()
    warnings.filterwarnings('ignore')
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    m = json.load(open(o.mappatura, encoding='utf-8-sig'))
    ws = openpyxl.load_workbook(o.export, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    Y = [dict(zip(rows[0], r)) for r in rows[1:]]
    g = lambda y, k: '' if y.get(k) is None else str(y[k]).strip()
    yext = collections.defaultdict(list)
    for y in Y:
        u = re.sub(r'^https?://(www\.)?kosgroup\.com/', '', g(y, 'URL del sito web > URL')).split('?')[0].strip('/')
        if u:
            yext[u].append(y)
    data = re.search(r'(\d{4}-\d{2}-\d{2})', os.path.basename(o.export))
    data = data.group(1) if data else ''

    def servizio(x):
        return x.get('sezione') == 'servizi-e-cure' and not x.get('mancante')

    def sezione(uri):
        p = (uri or '').strip('/').split('/')
        s = SEZIONI.get(p[2] if len(p) > 2 else '', 'Altro')
        return s + (' — trattamenti' if 'trattamenti' in p else '')

    righe = []
    for s in m['schede']:
        A = {x['id']: x for x in s['dalla_scheda'] if servizio(x)}
        B = {x['id']: x for x in s['dal_servizio'] if servizio(x)}
        pagine = {**A, **B}
        per_nome = {norma(x['pagetitle']): x for x in pagine.values()}
        ent = yext.get(s['uri'].strip('/'), [])
        voci = [v.strip() for e in ent for v in g(e, 'Servizi e Cure').split(',') if v.strip()]
        C, extra = set(), []
        for v in dict.fromkeys(voci):
            x = per_nome.get(norma(ALIAS.get(v, v)))
            if x:
                C.add(x['id'])
            else:
                extra.append(v)
        base = [s['pagetitle'], 'https://kosgroup.com/' + s['uri'].strip('/'),
                ', '.join(g(e, 'ID entità') for e in ent) or '(nessuna entità Yext)']
        for pid, x in sorted(pagine.items(), key=lambda kv: kv[1]['uri']):
            a_, b_, c_ = pid in A, pid in B, pid in C
            if a_ and b_ and c_:
                stato, prop = 'Coerente', 'Confermare'
            elif a_ and b_ and 'trattamenti' in x['uri']:
                stato, prop = 'Trattamento', 'Confermare (i trattamenti non sono tra i filtri di Yext)'
            elif a_ and b_:
                stato, prop = 'Solo sul sito', ('Aggiungere in Yext' if ent else 'Creare l\'entità in Yext')
            elif a_:
                # sulla scheda (fonte, curata nel rifacimento 2025) ma non sulla pagina del servizio
                tr = 'trattamenti' in x['uri']
                stato = 'Sulla scheda'
                prop = 'Confermata dalla scheda: si completa la pagina del servizio' + (
                    '' if tr or c_ else (' e si aggiunge in Yext' if ent else ''))
            else:
                # solo sulla pagina del servizio (in gran parte copiate dai filtri Yext): da verificare
                stato = 'Da verificare'
                prop = 'La struttura offre questo servizio? Presente solo sulla pagina del servizio' + (
                    ' e nei filtri Yext' if c_ else '')
            righe.append(base + [x['pagetitle'], sezione(x['uri']), 'https://kosgroup.com/' + x['uri'].strip('/'),
                                 'sì' if a_ else '', 'sì' if b_ else '', 'sì' if c_ else '', stato, prop, '', ''])
        for v in extra:
            ali = ALIAS.get(v, v)
            righe.append(base + [v, '', '', '', '', 'sì', 'Solo in Yext',
                                 'Nessuna pagina del sito associata con questo nome: associare la pagina sul sito o '
                                 'togliere la voce da Yext' + (f' (nome corretto «{ali}»)' if ali != v else ''), '', ''])

    wb = openpyxl.Workbook()
    F = 'Arial'
    intest = ['Struttura', 'Scheda sul sito', 'Entità Yext', 'Servizio / cura', 'Sezione del sito', 'Pagina del servizio',
              'Sito: dalla scheda', 'Sito: dalla pagina del servizio', 'Yext: Servizi e Cure', 'Stato', 'Proposta',
              'Esito KOS', 'Note KOS']
    c = wb.active
    c.title = 'Confronto'
    c.append(intest)
    for r in righe:
        c.append(r)
    larg = [34, 40, 20, 40, 30, 50, 12, 14, 12, 16, 56, 16, 36]
    for i, w in enumerate(larg, 1):
        c.column_dimensions[get_column_letter(i)].width = w
    giallo = PatternFill('solid', fgColor='FFF2CC')
    colori = {'Coerente': 'E2EFDA', 'Trattamento': 'EDEDED', 'Sulla scheda': 'E2EFDA', 'Da verificare': 'F8CBAD', 'Solo sul sito': 'DDEBF7', 'Solo in Yext': 'FCE4D6', 'Sito incompleto': 'FCE4D6',
              'Sito incoerente': 'F8CBAD'}
    for row in c.iter_rows(min_row=1, max_row=c.max_row):
        for cell in row:
            cell.font = Font(name=F, size=10, bold=cell.row == 1)
            cell.alignment = Alignment(vertical='top', wrap_text=cell.column in (11, 13))
        if row[0].row > 1:
            row[9].fill = PatternFill('solid', fgColor=colori.get(row[9].value, 'FFFFFF'))
            row[11].fill = giallo
            row[12].fill = giallo
    c.freeze_panes = 'D2'
    c.auto_filter.ref = f'A1:{get_column_letter(len(intest))}{c.max_row}'
    dv = DataValidation(type='list', formula1='"Confermato,Da aggiungere,Da togliere,Da correggere"', allow_blank=True)
    c.add_data_validation(dv)
    dv.add(f'L2:L{c.max_row}')

    r = wb.create_sheet('Riepilogo')
    r.append(['Struttura', 'Associazioni', 'Coerente', 'Trattamento', 'Solo sul sito', 'Sulla scheda',
              'Da verificare', 'Solo in Yext', 'Con esito KOS'])
    strutture = list(dict.fromkeys(x[0] for x in righe))
    n = c.max_row
    for i, s in enumerate(strutture, 2):
        r.cell(i, 1, s)
        r.cell(i, 2, f'=COUNTIF(Confronto!$A$2:$A${n},A{i})')
        for j, st in enumerate(['Coerente', 'Trattamento', 'Solo sul sito', 'Sulla scheda', 'Da verificare', 'Solo in Yext'], 3):
            r.cell(i, j, f'=COUNTIFS(Confronto!$A$2:$A${n},$A{i},Confronto!$J$2:$J${n},"{st}")')
        r.cell(i, 9, f'=COUNTIFS(Confronto!$A$2:$A${n},$A{i},Confronto!$L$2:$L${n},"<>")')
    t = len(strutture) + 2
    r.cell(t, 1, 'Totale')
    for j in range(2, 10):
        col = get_column_letter(j)
        r.cell(t, j, f'=SUM({col}2:{col}{t - 1})')
    for row in r.iter_rows(min_row=1, max_row=t):
        for cell in row:
            cell.font = Font(name=F, size=10, bold=cell.row in (1, t))
    r.column_dimensions['A'].width = 52
    for j in range(2, 10):
        r.column_dimensions[get_column_letter(j)].width = 15
    r.freeze_panes = 'B2'

    lg = wb.create_sheet('Legenda', 0)
    testo = [
        ('Confronto delle associazioni tra strutture e servizi e cure', True),
        (f'Sito kosgroup.com (produzione) e filtri «Servizi e Cure» di Yext (export del {data}).', False),
        ('', False),
        ('Fonti confrontate', True),
        ('Sito: dalla scheda — la scheda della struttura indica il servizio (campi di mappatura della scheda).', False),
        ('Sito: dalla pagina del servizio — la pagina del servizio indica la struttura (elenco strutture della pagina).', False),
        ('Yext: Servizi e Cure — il servizio è tra i filtri della mappa Yext per quella struttura.', False),
        ('', False),
        ('Stato', True),
        ('Coerente — presente in tutte e tre le fonti.', False),
        ('Trattamento — trattamento associato sul sito in entrambe le direzioni; Yext non ha i trattamenti tra i filtri.', False),
        ('Solo sul sito — associazione presente sul sito (in entrambe le direzioni) ma non in Yext.', False),
        ('Sulla scheda — indicata dalla scheda della struttura (fonte) ma non dalla pagina del servizio: si completa.', False),
        ('Da verificare — indicata solo dalla pagina del servizio (spesso copiata dai filtri Yext): KOS conferma o no.', False),
        ('Solo in Yext — voce di Yext senza pagina del sito associata alla struttura.', False),
        ('', False),
        ('Cosa compilare', True),
        ('Nel foglio «Confronto», colonne gialle: «Esito KOS» (Confermato, Da aggiungere, Da togliere, Da correggere)', False),
        ('e, se serve, «Note KOS». Esempio: Villa dei Pini Civitanova — Risonanza magnetica — Confermato.', False),
        ('Le associazioni confermate diventano la fonte unica: sito, Yext/RetailTune e dato strutturato le useranno uguali.', False),
    ]
    for i, (tx, b) in enumerate(testo, 1):
        cell = lg.cell(i, 1, tx)
        cell.font = Font(name=F, size=12 if i == 1 else 10, bold=b)
    lg.column_dimensions['A'].width = 120

    os.makedirs(o.uscita, exist_ok=True)
    out = os.path.join(o.uscita, f'confronto-servizi-{data or "kos"}.xlsx')
    wb.save(out)
    st = collections.Counter(x[9] for x in righe)
    print(f'{len(righe)} righe, {len(strutture)} strutture → {out}')
    print(dict(st))


if __name__ == '__main__':
    main()
